"""End-to-end tests for the complete pipeline stack.

Uses synthetic data (no real CIF required) plus the existing
Sample Output CSVs committed to the repository.

Tests:
  1. test_sample_output_csv_structure      — structural checks on Sample Output CSVs
  2. test_result_store                     — save / load / list / delete cycle
  3. test_electric_pipeline_synthetic      — electric energy pipeline with synthetic inputs
  4. test_solar_pipeline_synthetic         — solar analysis pipeline with synthetic inputs
  5. test_electric_transform               — format transform from scraper → pipeline columns
  6. test_combine_route_csv_deduplication  — multi-result route deduplication
  7. test_combine_timetable_deduplication  — service-level timetable merge / dedup
"""

from __future__ import annotations

import io
import struct
import tempfile
import textwrap
from pathlib import Path
from datetime import datetime, timedelta

import pandas as pd
import pytest

# Project root relative to this file
_REPO = Path(__file__).resolve().parents[3]
_SAMPLE_DIR = _REPO / "backend" / "data" / "sample"


# ---------------------------------------------------------------------------
# 1. Sample CSV structure (uses backend/data/sample/)
# ---------------------------------------------------------------------------

class TestSampleOutputStructure:
    """Validate that the bundled sample CSVs have the expected shape."""

    TIMETABLE_REQUIRED_COLS = {"route_variant", "date", "departure_time"}
    ROUTE_REQUIRED_COLS = {
        "route_variant", "seq", "from_station", "to_station",
        "run_min", "wait_min",
    }

    def test_sample_timetable_columns(self):
        df = pd.read_csv(_SAMPLE_DIR / "timetable_sample.csv")
        # timetable_sample uses scraper column names (departure_time instead of dep_time)
        assert "route_variant" in df.columns or "train_route" in df.columns, (
            "Sample timetable missing a route column"
        )
        assert "date" in df.columns

    def test_sample_route_columns(self):
        df = pd.read_csv(_SAMPLE_DIR / "route_sample.csv")
        assert self.ROUTE_REQUIRED_COLS.issubset(set(df.columns)), (
            f"Sample route missing columns: {self.ROUTE_REQUIRED_COLS - set(df.columns)}"
        )

    def test_sample_timetable_nonempty(self):
        df = pd.read_csv(_SAMPLE_DIR / "timetable_sample.csv")
        assert len(df) > 0

    def test_sample_route_seq_monotonic_per_variant(self):
        df = pd.read_csv(_SAMPLE_DIR / "route_sample.csv")
        for variant, grp in df.groupby("route_variant"):
            seqs = grp["seq"].tolist()
            assert seqs == sorted(seqs), (
                f"Route variant '{variant}' seq not monotonic: {seqs}"
            )

    def test_sample_route_has_at_least_one_variant(self):
        df = pd.read_csv(_SAMPLE_DIR / "route_sample.csv")
        assert df["route_variant"].nunique() >= 1


# ---------------------------------------------------------------------------
# 2. ResultStore
# ---------------------------------------------------------------------------

class TestResultStore:
    """Save / load / list / delete cycle."""

    def test_save_and_load(self, tmp_path):
        from backend.app.services.result_store import ResultStore
        store = ResultStore(str(tmp_path))
        store.save("id1", "tt_csv", "rt_csv", "dbg_csv", {"station": "FNB"})
        assert store.load_csv("id1", "timetable") == "tt_csv"
        assert store.load_csv("id1", "route") == "rt_csv"
        assert store.load_csv("id1", "debug") == "dbg_csv"

    def test_metadata_round_trip(self, tmp_path):
        from backend.app.services.result_store import ResultStore
        store = ResultStore(str(tmp_path))
        meta = {"station": "FNB", "timetable_rows": 42}
        store.save("id2", "a", "b", "c", meta)
        loaded = store.load_metadata("id2")
        assert loaded["station"] == "FNB"
        assert loaded["timetable_rows"] == 42

    def test_list_results(self, tmp_path):
        from backend.app.services.result_store import ResultStore
        store = ResultStore(str(tmp_path))
        store.save("a1", "t", "r", "d", {"station": "A"})
        store.save("b2", "t", "r", "d", {"station": "B"})
        results = store.list_results()
        ids = {r["generation_id"] for r in results}
        assert {"a1", "b2"}.issubset(ids)

    def test_delete(self, tmp_path):
        from backend.app.services.result_store import ResultStore
        store = ResultStore(str(tmp_path))
        store.save("x1", "t", "r", "d", {})
        assert store.exists("x1")
        store.delete("x1")
        assert not store.exists("x1")
        assert store.load_csv("x1", "timetable") is None

    def test_write_csv_updates_on_disk(self, tmp_path):
        from backend.app.services.result_store import ResultStore
        store = ResultStore(str(tmp_path))
        store.save("e1", "original", "r", "d", {})
        store.write_csv("e1", "timetable", "edited")
        assert store.load_csv("e1", "timetable") == "edited"


# ---------------------------------------------------------------------------
# 3. Electric pipeline — synthetic inputs
# ---------------------------------------------------------------------------

def _make_timetable_csv(scraper_format: bool = True) -> str:
    """Return a minimal timetable CSV."""
    if scraper_format:
        return textwrap.dedent("""\
            route_variant,departure_time,date,train_class,number_of_coaches
            R1,08:00:00,2025-01-15,Class377,8
            R1,09:00:00,2025-01-15,Class377,8
        """)
    return textwrap.dedent("""\
        route_variant,dep_time,date,train_type,cars
        R1,08:00,15/01/2025,Class377,8
        R1,09:00,15/01/2025,Class377,8
    """)


def _make_route_csv() -> str:
    return textwrap.dedent("""\
        route_variant,seq,from_station,to_station,distance_miles,run_min,wait_min
        R1,1,Station A,Station B,10.0,15,2
        R1,2,Station B,Station C,8.0,12,1
    """)


def _make_rolling_stock_csv(tmp_path: Path) -> Path:
    path = tmp_path / "rolling_stock.csv"
    path.write_text(textwrap.dedent("""\
        train_type,kwh_per_km_per_car,aux_kw_per_car,drive_eff,regen_eff,line_losses_pct
        Class377,2.0,10.0,0.9,0.0,0.0
    """))
    return path


def _make_station_points_csv(tmp_path: Path) -> Path:
    path = tmp_path / "station_points.csv"
    path.write_text(textwrap.dedent("""\
        Station,TSS
        Station A,TSS_NORTH
        Station B,TSS_NORTH
        Station C,TSS_SOUTH
    """))
    return path


class TestElectricPipeline:

    def test_returns_tss_dict(self, tmp_path):
        from backend.app.services.electric_pipeline import run_pipeline
        rs = _make_rolling_stock_csv(tmp_path)
        sp = _make_station_points_csv(tmp_path)
        result, _debug = run_pipeline(_make_timetable_csv(), _make_route_csv(), rs, sp)
        assert isinstance(result, dict)
        assert len(result) > 0, "Pipeline returned no TSS outputs"

    def test_tss_names_are_safe_strings(self, tmp_path):
        from backend.app.services.electric_pipeline import run_pipeline
        rs = _make_rolling_stock_csv(tmp_path)
        sp = _make_station_points_csv(tmp_path)
        result, _debug = run_pipeline(_make_timetable_csv(), _make_route_csv(), rs, sp)
        for name in result:
            assert name.replace("_", "").isalnum() or "_" in name, (
                f"TSS name not safe: {name!r}"
            )

    def test_each_tss_csv_has_48_bin_columns(self, tmp_path):
        from backend.app.services.electric_pipeline import run_pipeline, BIN_LABELS
        rs = _make_rolling_stock_csv(tmp_path)
        sp = _make_station_points_csv(tmp_path)
        result, _debug = run_pipeline(_make_timetable_csv(), _make_route_csv(), rs, sp)
        for tss_name, csv_text in result.items():
            df = pd.read_csv(io.StringIO(csv_text))
            bin_cols = [c for c in df.columns if c in BIN_LABELS]
            assert len(bin_cols) == 48, (
                f"{tss_name}: expected 48 bin columns, got {len(bin_cols)}"
            )

    def test_bin_values_are_non_negative(self, tmp_path):
        from backend.app.services.electric_pipeline import run_pipeline, BIN_LABELS
        rs = _make_rolling_stock_csv(tmp_path)
        sp = _make_station_points_csv(tmp_path)
        result, _debug = run_pipeline(_make_timetable_csv(), _make_route_csv(), rs, sp)
        for tss_name, csv_text in result.items():
            df = pd.read_csv(io.StringIO(csv_text))
            bin_cols = [c for c in df.columns if c in BIN_LABELS]
            assert (df[bin_cols].fillna(0) >= 0).all().all(), (
                f"{tss_name}: negative energy values found"
            )

    def test_bin_values_no_nan(self, tmp_path):
        from backend.app.services.electric_pipeline import run_pipeline, BIN_LABELS
        rs = _make_rolling_stock_csv(tmp_path)
        sp = _make_station_points_csv(tmp_path)
        result, _debug = run_pipeline(_make_timetable_csv(), _make_route_csv(), rs, sp)
        for tss_name, csv_text in result.items():
            df = pd.read_csv(io.StringIO(csv_text))
            bin_cols = [c for c in df.columns if c in BIN_LABELS]
            assert not df[bin_cols].isna().any().any(), (
                f"{tss_name}: NaN values in bin columns"
            )

    def test_total_units_equals_sum_of_bins(self, tmp_path):
        from backend.app.services.electric_pipeline import run_pipeline, BIN_LABELS
        rs = _make_rolling_stock_csv(tmp_path)
        sp = _make_station_points_csv(tmp_path)
        result, _debug = run_pipeline(_make_timetable_csv(), _make_route_csv(), rs, sp)
        for tss_name, csv_text in result.items():
            df = pd.read_csv(io.StringIO(csv_text))
            bin_cols = [c for c in df.columns if c in BIN_LABELS]
            computed = df[bin_cols].sum(axis=1).round(6)
            stored = df["Total Units"].round(6)
            pd.testing.assert_series_equal(
                computed.reset_index(drop=True),
                stored.reset_index(drop=True),
                check_names=False,
            )


class TestDwellLogic:
    """stop_type-aware dwell time and default-dwell behaviour."""

    def _run_segments(self, route_csv: str, tmp_path) -> list[dict]:
        """Run _expand_services on a single-service timetable, return all event records."""
        from backend.app.services.electric_pipeline import _expand_services, _load_energy_params
        import pandas as pd
        tt = pd.read_csv(io.StringIO(textwrap.dedent("""\
            route_variant,dep_time,date,train_type,cars
            R1,08:00,15/01/2025,Class377,4
        """)))
        route = pd.read_csv(io.StringIO(route_csv))
        stations = pd.read_csv(io.StringIO(textwrap.dedent("""\
            Station,TSS
            A,TSS1
            B,TSS1
            C,TSS1
            D,TSS1
        """)))
        ep = pd.DataFrame([{
            "train_type": "Class377", "kwh_per_km_per_car": 1.0,
            "aux_kw_per_car": 6.0, "drive_eff": 1.0,
            "regen_eff": 0.0, "line_losses_pct": 0.0,
        }])
        events = _expand_services(tt, route, stations, ep)
        return events[events.get("kind", pd.Series()).notna()].to_dict("records") \
            if "kind" in events.columns else []

    def test_pass_through_station_has_zero_dwell(self, tmp_path):
        """stop_type='pass' → no dwell event, even when wait_min=2.0."""
        route = textwrap.dedent("""\
            route_variant,seq,from_station,to_station,stop_type,distance,run_min,wait_min
            R1,1,A,B,pass,0.5,2,2.0
        """)
        events = self._run_segments(route, tmp_path)
        dwell_events = [e for e in events if e.get("kind") == "dwell"]
        assert dwell_events == [], (
            f"Expected no dwell for pass-through, got: {dwell_events}"
        )

    def test_stop_with_blank_wait_defaults_to_half_minute(self, tmp_path):
        """stop_type='stop' with blank wait_min → dwell of 0.5 min (aux energy > 0)."""
        route = textwrap.dedent("""\
            route_variant,seq,from_station,to_station,stop_type,distance,run_min,wait_min
            R1,1,A,B,stop,0.5,2,
        """)
        events = self._run_segments(route, tmp_path)
        dwell_events = [e for e in events if e.get("kind") == "dwell"]
        assert len(dwell_events) == 1, f"Expected 1 dwell event, got {len(dwell_events)}"
        # dwell duration = 0.5 min → aux_kwh_dwell = 6.0 × 4 × (0.5/60) = 0.2 kWh
        expected_kwh = 6.0 * 4 * (0.5 / 60)
        assert abs(dwell_events[0]["kwh"] - expected_kwh) < 1e-6, (
            f"Expected {expected_kwh:.6f} kWh, got {dwell_events[0]['kwh']:.6f}"
        )

    def test_stop_with_zero_wait_defaults_to_half_minute(self, tmp_path):
        """stop_type='stop' with explicit wait_min=0 → treated as blank, defaults to 0.5."""
        route = textwrap.dedent("""\
            route_variant,seq,from_station,to_station,stop_type,distance,run_min,wait_min
            R1,1,A,B,stop,0.5,2,0
        """)
        events = self._run_segments(route, tmp_path)
        dwell_events = [e for e in events if e.get("kind") == "dwell"]
        assert len(dwell_events) == 1
        expected_kwh = 6.0 * 4 * (0.5 / 60)
        assert abs(dwell_events[0]["kwh"] - expected_kwh) < 1e-6

    def test_stop_with_explicit_wait_min_used(self, tmp_path):
        """stop_type='stop' with wait_min=2.0 → uses 2.0, not the 0.5 default."""
        route = textwrap.dedent("""\
            route_variant,seq,from_station,to_station,stop_type,distance,run_min,wait_min
            R1,1,A,B,stop,0.5,2,2.0
        """)
        events = self._run_segments(route, tmp_path)
        dwell_events = [e for e in events if e.get("kind") == "dwell"]
        assert len(dwell_events) == 1
        expected_kwh = 6.0 * 4 * (2.0 / 60)
        assert abs(dwell_events[0]["kwh"] - expected_kwh) < 1e-6, (
            f"Expected {expected_kwh:.6f} kWh, got {dwell_events[0]['kwh']:.6f}"
        )

    def test_no_stop_type_column_behaves_as_stop(self, tmp_path):
        """Route CSV without stop_type column: applies stop default (0.5 min for blank wait)."""
        route = textwrap.dedent("""\
            route_variant,seq,from_station,to_station,distance,run_min,wait_min
            R1,1,A,B,0.5,2,
        """)
        events = self._run_segments(route, tmp_path)
        dwell_events = [e for e in events if e.get("kind") == "dwell"]
        assert len(dwell_events) == 1
        expected_kwh = 6.0 * 4 * (0.5 / 60)
        assert abs(dwell_events[0]["kwh"] - expected_kwh) < 1e-6


class TestElectricTransform:
    """format transform: scraper columns → pipeline columns."""

    def test_scraper_columns_accepted(self, tmp_path):
        from backend.app.services.electric_pipeline import run_pipeline
        rs = _make_rolling_stock_csv(tmp_path)
        sp = _make_station_points_csv(tmp_path)
        # Should not raise even with scraper-format column names
        result, _debug = run_pipeline(
            _make_timetable_csv(scraper_format=True),
            _make_route_csv(), rs, sp,
        )
        assert len(result) > 0

    def test_native_columns_accepted(self, tmp_path):
        from backend.app.services.electric_pipeline import run_pipeline
        rs = _make_rolling_stock_csv(tmp_path)
        sp = _make_station_points_csv(tmp_path)
        result, _debug = run_pipeline(
            _make_timetable_csv(scraper_format=False),
            _make_route_csv(), rs, sp,
        )
        assert len(result) > 0


class TestCombineRoutes:
    """Multi-result route deduplication."""

    def test_combine_route_deduplicates(self):
        from backend.app.services.electric_pipeline import combine_route_csvs
        csv_a = textwrap.dedent("""\
            route_variant,seq,from_station,to_station,distance_miles,run_min,wait_min
            R1,1,A,B,10,15,2
            R1,2,B,C,8,12,1
        """)
        # Identical to csv_a — duplicate should be removed
        csv_b = textwrap.dedent("""\
            route_variant,seq,from_station,to_station,distance_miles,run_min,wait_min
            R1,1,A,B,10,15,2
            R1,2,B,C,8,12,1
        """)
        combined = combine_route_csvs([csv_a, csv_b])
        df = pd.read_csv(io.StringIO(combined))
        assert len(df) == 2, f"Expected 2 unique rows, got {len(df)}"

    def test_combine_route_keeps_distinct_variants(self):
        from backend.app.services.electric_pipeline import combine_route_csvs
        csv_a = "route_variant,seq,from_station,to_station,distance_miles,run_min,wait_min\nR1,1,A,B,10,15,2\n"
        csv_b = "route_variant,seq,from_station,to_station,distance_miles,run_min,wait_min\nR2,1,X,Y,5,8,0\n"
        combined = combine_route_csvs([csv_a, csv_b])
        df = pd.read_csv(io.StringIO(combined))
        assert set(df["route_variant"]) == {"R1", "R2"}


# ---------------------------------------------------------------------------
# 7. Timetable combine / service-level deduplication
# ---------------------------------------------------------------------------

class TestCombineTimetable:
    """Timetable CSV merge and service-level deduplication tests.

    Dedup key: (date, train_uid, origin_departure).
    Fallback:  (date, departure_time, route_variant) for legacy CSVs.
    """

    _COLS = ("route_variant,train_uid,origin_departure,stop_type,"
             "date,departure_time,train_class,number_of_coaches")

    def _csv(self, rows: list[str]) -> str:
        return self._COLS + "\n" + "\n".join(rows) + "\n"

    def test_exact_duplicate_removed(self):
        """Identical row present in two CSVs → 1 row after merge."""
        from backend.app.services.electric_pipeline import combine_csvs
        row = "Waterloo - Basingstoke,W12345,0755,stop,2025-01-06,08:05:00,Electric,8"
        merged = combine_csvs([self._csv([row]), self._csv([row])])
        df = pd.read_csv(io.StringIO(merged))
        assert len(df) == 1, f"Expected 1 row, got {len(df)}"

    def test_same_service_different_stations_deduped(self):
        """Same physical service queried at two different stations (different
        departure_time, same train_uid + origin_departure) → 1 row after merge,
        keeping the earliest departure_time."""
        from backend.app.services.electric_pipeline import combine_csvs
        fleet_row = "Waterloo - Basingstoke,W12345,0755,stop,2025-01-06,08:05:00,Electric,8"
        camb_row  = "Waterloo - Basingstoke,W12345,0755,stop,2025-01-06,08:15:00,Electric,8"
        merged = combine_csvs([self._csv([fleet_row]), self._csv([camb_row])])
        df = pd.read_csv(io.StringIO(merged))
        assert len(df) == 1, f"Expected 1 row, got {len(df)}"
        assert df.iloc[0]["departure_time"] == "08:05:00", (
            f"Expected earliest departure 08:05:00, got {df.iloc[0]['departure_time']}"
        )

    def test_stp_pn_same_uid_kept_separate(self):
        """STP P and N services sharing the same train_uid but running at different
        origin departure times must NOT be collapsed — they are distinct trains."""
        from backend.app.services.electric_pipeline import combine_csvs
        # P service: origin=0800; N service: origin=0900 (same UID, same date)
        p_row = "Waterloo - Basingstoke,A12345,0800,stop,2025-01-06,08:10:00,Electric,8"
        n_row = "Waterloo - Basingstoke,A12345,0900,stop,2025-01-06,09:10:00,Electric,8"
        merged = combine_csvs([self._csv([p_row, n_row])])
        df = pd.read_csv(io.StringIO(merged))
        assert len(df) == 2, (
            f"Expected 2 rows (P and N must stay separate), got {len(df)}"
        )

    def test_richer_row_wins_on_score_tie(self):
        """When (train_uid, origin_departure, date) match but one row has more
        data (train_class populated), the richer row is kept."""
        from backend.app.services.electric_pipeline import combine_csvs
        sparse = "Waterloo - Basingstoke,W12345,0755,stop,2025-01-06,08:05:00,,8"
        rich   = "Waterloo - Basingstoke,W12345,0755,stop,2025-01-06,08:15:00,Electric,8"
        merged = combine_csvs([self._csv([sparse]), self._csv([rich])])
        df = pd.read_csv(io.StringIO(merged))
        assert len(df) == 1, f"Expected 1 row, got {len(df)}"
        assert df.iloc[0]["train_class"] == "Electric", (
            f"Expected richer row (train_class=Electric) to win, got {df.iloc[0]['train_class']}"
        )

    def test_distinct_services_both_kept(self):
        """Two different services (different train_uids) on the same date must
        both be present after merge."""
        from backend.app.services.electric_pipeline import combine_csvs
        svc_x = "Waterloo - Basingstoke,W12345,0755,stop,2025-01-06,08:05:00,Electric,8"
        svc_y = "Waterloo - Basingstoke,W67890,0900,stop,2025-01-06,09:05:00,Electric,8"
        merged = combine_csvs([self._csv([svc_x]), self._csv([svc_y])])
        df = pd.read_csv(io.StringIO(merged))
        assert len(df) == 2, f"Expected 2 rows (X and Y), got {len(df)}"
        assert set(df["train_uid"]) == {"W12345", "W67890"}

    def test_fallback_to_legacy_key_without_train_uid(self):
        """Legacy CSVs that lack train_uid/origin_departure fall back to the
        (date, departure_time, route_variant) key and still deduplicate."""
        from backend.app.services.electric_pipeline import combine_csvs
        legacy_cols = "route_variant,stop_type,date,departure_time,train_class,number_of_coaches"
        row = "Waterloo - Basingstoke,stop,2025-01-06,08:05:00,Electric,8"
        csv_text = legacy_cols + "\n" + row + "\n"
        merged = combine_csvs([csv_text, csv_text])
        df = pd.read_csv(io.StringIO(merged))
        assert len(df) == 1, f"Expected 1 row after legacy dedup, got {len(df)}"


# ---------------------------------------------------------------------------
# 4. Solar pipeline — synthetic inputs
# ---------------------------------------------------------------------------

def _make_demand_csv() -> str:
    """Synthetic half-hour demand CSV (48 bins) for 3 days."""
    from backend.app.services.electric_pipeline import BIN_LABELS
    header = "Date,Day,Total Units," + ",".join(BIN_LABELS)
    rows = []
    for i in range(3):
        day = datetime(2025, 6, 1) + timedelta(days=i)
        # Simulate demand peaking at midday
        bins = [max(0.0, 10.0 - abs(j - 24) * 0.5) for j in range(48)]
        total = sum(bins)
        row = f"{day.strftime('%d/%m/%Y')},{day.strftime('%A')},{total:.4f},"
        row += ",".join(f"{v:.4f}" for v in bins)
        rows.append(row)
    return header + "\n" + "\n".join(rows) + "\n"


def _make_pvgis_csv() -> str:
    """Synthetic PVGIS CSV matching the expected format."""
    lines = [
        "Latitude: 51.295",
        "Longitude: -0.773",
        "Elevation: 65 m",
        "# some other metadata",
        "time,P,G(i),H_sun,T2m,WS10m,Int",
    ]
    # 3 days × 24 hours
    for i in range(3):
        day = datetime(2019, 6, 1) + timedelta(days=i)
        for h in range(24):
            dt_str = day.strftime("%Y%m%d") + f":{h:02d}00"
            # solar power peaks at midday
            p = max(0.0, 500.0 * (1 - abs(h - 12) / 6)) if 6 <= h <= 18 else 0.0
            lines.append(f"{dt_str},{p:.1f},0.0,0.0,15.0,3.0,0")
    return "\n".join(lines) + "\n"


class TestSolarPipeline:

    def test_returns_solar_result(self):
        from backend.app.services.solar_pipeline import run_solar_pipeline, SolarResult
        result = run_solar_pipeline(_make_demand_csv(), _make_pvgis_csv())
        assert isinstance(result, SolarResult)

    def test_avg_profile_png_is_valid_png(self):
        from backend.app.services.solar_pipeline import run_solar_pipeline
        result = run_solar_pipeline(_make_demand_csv(), _make_pvgis_csv())
        # PNG files start with the 8-byte PNG signature
        assert result.avg_profile_png[:8] == b"\x89PNG\r\n\x1a\n", (
            "avg_profile_png does not start with PNG signature"
        )

    def test_metrics_xlsx_is_nonempty(self):
        from backend.app.services.solar_pipeline import run_solar_pipeline
        result = run_solar_pipeline(_make_demand_csv(), _make_pvgis_csv())
        assert len(result.metrics_xlsx) > 0

    def test_demand_hourly_xlsx_is_nonempty(self):
        from backend.app.services.solar_pipeline import run_solar_pipeline
        result = run_solar_pipeline(_make_demand_csv(), _make_pvgis_csv())
        assert len(result.demand_hourly_xlsx) > 0

    def test_solar_share_pct_in_range(self):
        from backend.app.services.solar_pipeline import run_solar_pipeline
        result = run_solar_pipeline(_make_demand_csv(), _make_pvgis_csv())
        assert 0.0 <= result.solar_share_pct <= 100.0, (
            f"solar_share_pct out of range: {result.solar_share_pct}"
        )

    def test_utilisation_pct_in_range(self):
        from backend.app.services.solar_pipeline import run_solar_pipeline
        result = run_solar_pipeline(_make_demand_csv(), _make_pvgis_csv())
        assert 0.0 <= result.utilisation_pct <= 100.0, (
            f"utilisation_pct out of range: {result.utilisation_pct}"
        )

    def test_metrics_workbook_has_required_sheets(self):
        """Annual Summary, Seasonal Summary, and Matched Detail sheets must exist."""
        import openpyxl
        from backend.app.services.solar_pipeline import run_solar_pipeline
        result = run_solar_pipeline(_make_demand_csv(), _make_pvgis_csv())
        wb = openpyxl.load_workbook(io.BytesIO(result.metrics_xlsx))
        assert "Annual Summary" in wb.sheetnames
        assert "Seasonal Summary" in wb.sheetnames
        assert "Matched Detail" in wb.sheetnames

    def test_avg_profile_xlsx_has_24_hour_rows(self):
        import openpyxl
        from backend.app.services.solar_pipeline import run_solar_pipeline
        result = run_solar_pipeline(_make_demand_csv(), _make_pvgis_csv())
        wb = openpyxl.load_workbook(io.BytesIO(result.avg_profile_xlsx))
        ws = wb.active
        # 1 header + 24 data rows
        assert ws.max_row == 25, f"Expected 25 rows (1 header + 24 hours), got {ws.max_row}"
